"""Build the final project report (Word) from the EEE report template.

The template (read only) keeps its cover page, styles, table of contents,
headers/footers and numbering; the body is replaced with this project's
content. Equations are native Word equations (LaTeX -> MathML -> OMML).
Usage (.venv-phase10):  python src/build_report.py
Then src/finalise_report.ps1 updates the TOC/fields in Word and exports a PDF.
"""
from __future__ import annotations

import copy
import os
import re

import latex2mathml.converter
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from lxml import etree

from report_content import CONTENT, REFS

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TPL = r"D:\User\Downloads\EEE-xxx-project-report-template.docx"
OUT = os.path.join(ROOT, "report", "Group_13_EEE_402_G2_Project_Report.docx")
FIG = os.path.join(ROOT, "report", "figures")
EXPORT = os.path.join(ROOT, "report", "report_text_for_quillbot.txt")     # written on every build
HUMAN = os.path.join(ROOT, "report", "report_text_humanized.txt")         # used if present
PFIG = os.path.join(ROOT, "presentation", "figures")
XSL = etree.XSLT(etree.parse(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"))

TERM = "January 2026"
COURSE = "Artificial Intelligence and Machine Learning"
TITLE = "Telling Lung Adenocarcinoma from Squamous Cell Carcinoma on CT Using Hospital-Robust Radiomics"
SHORT = "CT-based ADC vs SCC Classification"
TEXT_W = 6.19                     # inches between the margins

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


# ======================================================================= helpers
def fill_nary(root):
    """MML2OMML leaves the summand of a sum empty (<m:e/>) when it is not one grouped term, and Word
    then draws an empty slot (a large gap) after the sign. Move the following terms into the sum,
    up to the next '=', ',', ';' or '≈' at the same level."""
    m = f"{{{M_NS}}}"
    stops = "=,;≈"
    for g in root.iter(m + "grow"):          # all sum signs the same size (no stretching around inner sums)
        g.set(m + "val", "0")
    for nary in list(root.iter(m + "nary")):
        e = nary.find(m + "e")
        if e is None or len(e):
            continue
        while True:
            sib = nary.getnext()
            if sib is None:
                break
            if sib.tag == m + "r":
                t = sib.find(m + "t")
                txt = t.text if t is not None and t.text else ""
                cut = min([txt.find(c) for c in stops if c in txt], default=-1)
                if cut == 0:
                    break
                if cut > 0:                       # split the run: head goes into the sum
                    head = copy.deepcopy(sib)
                    head.find(m + "t").text = txt[:cut]
                    t.text = txt[cut:]
                    e.append(head)
                    break
            e.append(sib)


def omml(latex, display=False):
    mml = latex2mathml.converter.convert(latex, display="block" if display else "inline")
    m = XSL(etree.fromstring(mml.encode())).getroot()
    fill_nary(m)
    if display:
        para = etree.Element(f"{{{M_NS}}}oMathPara")
        para.append(m)
        return para
    return m


def no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def set_runs(p, text):
    """put `text` in the first run (keeps its formatting), empty the others."""
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def strip_highlight(el):
    for h in el.iter(qn("w:highlight")):
        h.getparent().remove(h)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


class Numbers:
    """label -> number, assigned in a dry run (order of appearance)."""

    def __init__(self):
        self.fig, self.tab, self.eq, self.ref = {}, {}, {}, {}

    def resolve(self, text):
        def f(m):
            kind, keys = m.group(1), m.group(2)
            if kind == "cite":
                nums = []
                for k in keys.split(","):
                    k = k.strip()
                    if k not in self.ref:
                        self.ref[k] = len(self.ref) + 1
                    nums.append(self.ref[k])
                return "[" + ", ".join(str(n) for n in sorted(nums)) + "]"
            d = {"fig": self.fig, "tab": self.tab, "eq": self.eq}[kind]
            if keys not in d:
                return "??"
            return {"fig": "Figure ", "tab": "Table ", "eq": "Eq. ("}[kind] + str(d[keys]) + (")" if kind == "eq" else "")
        return re.sub(r"\{(fig|tab|eq|cite):([^}]+)\}", f, text)


# ====================================================================== builder
class Builder:
    def __init__(self, doc, nums, dry, overrides=None):
        self.doc, self.n, self.dry = doc, nums, dry
        self.body = doc.element.body
        self.nfig = self.ntab = self.neq = 0
        self.num_id = None
        self.uid = 0                     # running id of prose units (paragraphs, bullets, captions)
        self.units = []                  # (uid, resolved original text) for the QuillBot export
        self.overrides = overrides or {}  # uid -> humanized plain text
        self.rejected = []

    # ------------------------------------------------- humanized text units
    def _unit(self, text):
        """resolved text of one prose unit, replaced by its humanized version if one is available."""
        self.uid += 1
        orig = self.n.resolve(text)
        self.units.append((self.uid, orig))
        new = self.overrides.get(self.uid)
        if new is None:
            return orig
        ok, merged = merge_human(orig, new)
        if not ok:
            self.rejected.append((self.uid, merged))
            return orig
        return merged

    # ---------------------------------------------------------------- text
    def _rich(self, p, text, size=None, bold=False, italic=False, font=None):
        text = self.n.resolve(text)
        parts = re.split(r"(\$[^$]+\$)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("$") and part.endswith("$") and len(part) > 2:
                p._p.append(omml(part[1:-1]))
                continue
            for seg in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", part):
                if not seg:
                    continue
                b, i = bold, italic
                if seg.startswith("`") and seg.endswith("`") and len(seg) > 2:
                    r = p.add_run(seg[1:-1])
                    r.font.name = "Consolas"
                    r.font.size = Pt((size or 11) - 1.5)
                    continue
                if seg.startswith("**"):
                    seg, b = seg[2:-2], True
                elif seg.startswith("*") and seg.endswith("*") and len(seg) > 2:
                    seg, i = seg[1:-1], True
                r = p.add_run(seg)
                r.bold, r.italic = b or None, i or None
                if size:
                    r.font.size = Pt(size)
                if font:
                    r.font.name = font
                    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)

    def _para(self, style="Normal"):
        p = self.doc.add_paragraph(style=style)
        return p

    def h1(self, t):
        if not self.dry:
            self.doc.add_paragraph(t, style="Heading 1")

    def h2(self, t):
        if not self.dry:
            self.doc.add_paragraph(" " + t, style="Heading 2")

    def h3(self, t):
        if not self.dry:
            self.doc.add_paragraph(t, style="Heading 3")

    def sub(self, title):
        """small plain sub-heading on its own line (not numbered, not in the TOC)."""
        if self.dry:
            return
        p = self._para()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title)
        r.bold = True

    def p(self, text, align="justify", after=6, before=0, italic=False, keep=False, size=None, human=True):
        if self.dry:
            self.n.resolve(text)
            return
        if human:
            text = self._unit(text)
        p = self._para()
        p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        if keep:
            p.paragraph_format.keep_with_next = True
        self._rich(p, text, italic=italic, size=size)
        return p

    def lead(self, title, text):
        """paragraph that starts with a bold run-in title."""
        self.p(f"**{title}** {text}")

    def bullets(self, items, numbered=False):
        if self.dry:
            for it in items:
                self.n.resolve(it)
            return
        num_id = 3
        if numbered:
            num_id = self._new_numbering()
        for k, it in enumerate(items):
            p = self._para("List Paragraph")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(3 if k < len(items) - 1 else 6)
            pPr = p._p.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            il = OxmlElement("w:ilvl")
            il.set(qn("w:val"), "0")
            ni = OxmlElement("w:numId")
            ni.set(qn("w:val"), str(num_id))
            numPr.append(il)
            numPr.append(ni)
            pPr.insert(0, numPr) if pPr.find(qn("w:pStyle")) is None else pPr.find(qn("w:pStyle")).addnext(numPr)
            self._rich(p, self._unit(it))

    def _new_numbering(self):
        numbering = self.doc.part.numbering_part.element
        ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        nid = max(ids) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(nid))
        an = OxmlElement("w:abstractNumId")
        an.set(qn("w:val"), "2")
        num.append(an)
        lo = OxmlElement("w:lvlOverride")
        lo.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), "1")
        lo.append(so)
        num.append(lo)
        numbering.append(num)
        return nid

    # ------------------------------------------------------------ equations
    def eq(self, latex, label=None):
        self.neq += 1
        if label:
            self.n.eq[label] = self.neq
        if self.dry:
            return
        # display equation: borderless 3-column table [ blank | equation | (n) ]
        t = self.doc.add_table(rows=1, cols=3)
        no_borders(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for j, w in enumerate((0.55, TEXT_W - 1.1, 0.55)):
            t.cell(0, j).width = Inches(w)
            t.cell(0, j).vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        mid = t.cell(0, 1).paragraphs[0]
        mid.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mid.paragraph_format.space_before = Pt(2)
        mid.paragraph_format.space_after = Pt(2)
        mid._p.append(omml(latex, display=True))
        num = t.cell(0, 2).paragraphs[0]
        num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        num.add_run(f"({self.neq})")

    # -------------------------------------------------------------- figures
    def fig(self, path, caption, label, width=6.0):
        self.nfig += 1
        self.n.fig[label] = self.nfig
        if self.dry:
            self.n.resolve(caption)
            return
        full = path if os.path.isabs(path) else os.path.join(FIG if path.startswith("fig_") else PFIG, path)
        p = self._para()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(full, width=Inches(width))
        c = self._para("Caption")
        self._rich(c, f"Figure {self.nfig}: " + self._unit(caption))

    # --------------------------------------------------------------- tables
    def _cap_table(self, caption, human=True):
        c = self._para("Caption")
        c.paragraph_format.space_before = Pt(4)
        self._rich(c, f"Table {self.ntab}: " + (self._unit(caption) if human else caption))

    def table(self, header, rows, caption, label, widths=None, size=9.5, align=None, bold_last=False,
              shade_rows=None):
        self.ntab += 1
        self.n.tab[label] = self.ntab
        if self.dry:
            self.n.resolve(caption)
            for r in rows:
                for c in r:
                    self.n.resolve(str(c))
            return
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, row in enumerate([header] + rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(1)
                cp.paragraph_format.space_before = Pt(1)
                a = (align[j] if align else "l") if i else "c"
                cp.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
                                "r": WD_ALIGN_PARAGRAPH.RIGHT}[a]
                self._rich(cp, str(val), size=size, bold=(i == 0) or (bold_last and i == len(rows)))
                if i > 0 and shade_rows and (i - 1) in shade_rows:
                    shade(cell, "FFF2CC")
                if widths:
                    cell.width = Inches(widths[j])
        # repeat header row across pages
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        self._cap_table(caption)

    def code(self, left, right, caption, label):
        """template style: Consolas 7 pt, two columns."""
        self.ntab += 1
        self.n.tab[label] = self.ntab
        if self.dry:
            return
        t = self.doc.add_table(rows=1, cols=2)
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, txt in enumerate((left, right)):
            cell = t.cell(0, j)
            cell.width = Inches(TEXT_W / 2)
            shade(cell, "F2F2F2")
            first = True
            for line in txt.rstrip("\n").split("\n"):
                cp = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                cp.paragraph_format.space_after = Pt(0)
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.line_spacing = 1.0
                r = cp.add_run(line if line else " ")
                r.font.name = "Consolas"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                r._element.rPr.rFonts.set(qn("w:cs"), "Consolas")
                r.font.size = Pt(7)
        self._cap_table(caption, human=False)

    def logbook(self, rows, caption, label):
        """reuse the template's log-book table (header + one row)."""
        self.ntab += 1
        self.n.tab[label] = self.ntab
        if self.dry:
            return
        el = copy.deepcopy(self.proto_log)
        self.body.find(qn("w:sectPr")).addprevious(el)
        t = Table(el, self.doc._body)
        while len(t.rows) < len(rows) + 1:
            t.add_row()
        widths = [0.75, 2.75, 1.15, 0.75, 0.95]
        grid = el.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 1440)))
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Inches(w)
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                for extra in cell.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
                cp = cell.paragraphs[0]
                for r in list(cp.runs):
                    r._r.getparent().remove(r._r)
                cp.paragraph_format.space_after = Pt(1)
                self._rich(cp, val, size=9)
        self._cap_table(caption)

    def abstract(self, paras):
        """plain justified paragraphs (no shaded box); the keyword line is not humanized."""
        for txt in paras:
            self.p(txt, human=not txt.startswith("**Keywords:**"))

    def page_break(self):
        if not self.dry:
            self._para().add_run().add_break(__import__("docx").enum.text.WD_BREAK.PAGE)

    def references(self):
        if self.dry:
            return
        order = sorted(self.n.ref.items(), key=lambda kv: kv[1])
        for key, num in order:
            p = self._para()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.4)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))
            r = p.add_run(f"[{num}]\t")
            self._rich(p, REFS[key])


# ============================================================ QuillBot round trip
MATH = re.compile(r"\$[^$]+\$")
TOKEN = re.compile(r"\[\s*M\s*(\d+)\s*\]")
NUMBER = re.compile(r"\d+(?:[.,]\d+)*")


def plain_for_export(text):
    """resolved unit -> plain text: inline math becomes [M1], [M2]...; markup removed."""
    k = [0]

    def tok(_):
        k[0] += 1
        return f"[M{k[0]}]"
    out = MATH.sub(tok, text)
    out = re.sub(r"\*\*([^*]+)\*\*", r"\1", out)
    out = re.sub(r"\*([^*]+)\*", r"\1", out)
    return out.replace("`", "")


def merge_human(orig, new):
    """put the original inline math back into a humanized unit and check that nothing was lost.
    Returns (ok, text or reason)."""
    maths = MATH.findall(orig)
    new = " ".join(new.split())
    found = [int(x) for x in TOKEN.findall(new)]
    if sorted(found) != list(range(1, len(maths) + 1)):
        return False, f"math tokens {found} instead of 1..{len(maths)}"
    plain_orig = MATH.sub(" ", orig)
    plain_new = TOKEN.sub(" ", new)
    nums = lambda s: {n.replace(",", "") for n in NUMBER.findall(s)}  # noqa: E731  (1,041 == 1041)
    lost = sorted(nums(plain_orig) - nums(plain_new))
    if lost:
        return False, f"numbers missing: {lost}"
    out = TOKEN.sub(lambda m: maths[int(m.group(1)) - 1], new)
    out = re.sub(r"\bet al\.", "*et al.*", out)
    return True, out


def read_humanized(path):
    if not os.path.isfile(path):
        return {}
    txt = open(path, encoding="utf8").read()
    parts = re.split(r"\[\s*P\s*0*(\d+)\s*\]", txt)
    return {int(parts[i]): parts[i + 1].strip() for i in range(1, len(parts) - 1, 2) if parts[i + 1].strip()}


def write_export(units, path):
    with open(path, "w", encoding="utf8") as f:
        for uid, text in units:
            f.write(f"[P{uid:03d}]\n{plain_for_export(text)}\n\n")


# ================================================================ cover + frame
def edit_cover(doc):
    P = doc.paragraphs
    set_runs(P[3], f"EEE 402 ({TERM})")
    set_runs(P[4], COURSE)
    P[4].runs[1].text = " Laboratory"
    sec = P[6]
    for r in sec.runs:
        r.text = {"A1": "G2", "01": "13"}.get(r.text, r.text)
    if "G2" not in sec.text or "13" not in sec.text:
        set_runs(sec, "Section: G2 Group: 13")
    set_runs(P[8], TITLE)
    set_runs(P[15], "Tanvir Hossain, Lecturer")
    set_runs(P[16], "Anindya Bhattacharjee, Adjunct Lecturer")
    # academic honesty statement (text box: DrawingML + VML fallback)
    box = list(doc.element.body.iterchildren())[21]
    for tx in box.iter(qn("w:txbxContent")):
        tbl = tx.find(qn("w:tbl"))
        rows = tbl.findall(qn("w:tr"))
        for tr in rows[1:]:
            tbl.remove(tr)
        cells = rows[0].findall(qn("w:tc"))
        for wp in cells[0].findall(qn("w:p")):
            ts = list(wp.iter(qn("w:t")))
            txt = "".join(t.text or "" for t in ts)
            new = None
            if txt.startswith("Full Name"):
                new = "Full Name: Rahib Mahasin"
            elif txt.startswith("Student ID"):
                new = "Student ID: 2106119"
            if new and ts:
                ts[0].text = new
                for t in ts[1:]:
                    t.text = ""
        for c in cells[1:]:
            for t in c.iter(qn("w:t")):
                t.text = ""
    # one signature row was removed: shrink the box by 78 pt (DrawingML and VML copies)
    cut = 78 * 12700
    for tag in (qn("wp:extent"), "{http://schemas.openxmlformats.org/drawingml/2006/main}ext"):
        for e in box.iter(tag):
            if e.get("cy") == "2761615":
                e.set("cy", str(2761615 - cut))
    for e in box.iter("{urn:schemas-microsoft-com:vml}shape"):
        e.set("style", e.get("style", "").replace("height:217.45pt", "height:139.45pt"))
    strip_highlight(doc.element.body)


def edit_footers(doc):
    swap = {"Title of the Project": SHORT, "xxx": "402", "416": "402", "January 2022": TERM, "A1": "G2", "X": "13"}
    seen = set()
    for rel in doc.part.rels.values():
        if "footer" in rel.reltype or "header" in rel.reltype:
            part = rel.target_part
            if id(part) in seen:
                continue
            seen.add(id(part))
            el = part.element
            for t in el.iter(qn("w:t")):
                if t.text in swap:
                    t.text = swap[t.text]
            strip_highlight(el)


def main():
    doc = Document(TPL)
    body = doc.element.body
    kids = list(body.iterchildren())
    # prototypes from the template body
    proto_abs = copy.deepcopy(kids[28])
    proto_log = copy.deepcopy(kids[97])
    # drop everything after the TOC section break (keep the final sectPr)
    for el in kids[27:]:
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    edit_cover(doc)
    edit_footers(doc)

    nums = Numbers()
    overrides = read_humanized(HUMAN)
    for dry in (True, False):
        b = Builder(doc, nums, dry, overrides)
        b.proto_abs, b.proto_log = proto_abs, proto_log
        CONTENT(b)
    write_export(b.units, EXPORT)
    print(f"prose units: {len(b.units)} -> {EXPORT}")
    if overrides:
        used = len([u for u, _ in b.units if u in overrides]) - len(b.rejected)
        print(f"humanized text applied to {used} of {len(b.units)} units")
        for uid, why in b.rejected:
            print(f"  [P{uid:03d}] kept original: {why}")
        extra = sorted(set(overrides) - {u for u, _ in b.units})
        if extra:
            print("  unknown ids in humanized file:", extra)
    # sanity: every reference key exists, no unresolved labels
    missing = [k for k in nums.ref if k not in REFS]
    assert not missing, missing
    unused = [k for k in REFS if k not in nums.ref]
    if unused:
        print("WARNING: references never cited:", unused)
    xml = etree.tostring(body, encoding="unicode")
    assert "??" not in re.sub(r"<[^>]+>", "", xml), "unresolved cross-reference"
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("saved", OUT, "| figures", len(nums.fig), "| tables", len(nums.tab), "| equations", len(nums.eq),
          "| references", len(nums.ref))


if __name__ == "__main__":
    main()
