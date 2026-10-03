"""Build the final course presentation on top of the EEE department template.

The template file is only read; output -> presentation/ADC_SCC_Final_Presentation.pptx
Usage (.venv-phase10):  python src/build_slides.py
"""
from __future__ import annotations

import copy
import os
import re

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TEMPLATE = r"D:/User/Downloads/EEExxx-Final-Project-FinalDemo-Template.pptx"
FIG = os.path.join(ROOT, "presentation", "figures")
OUT = os.path.join(ROOT, "presentation", "ADC_SCC_Final_Presentation.pptx")

FOOT_LEFT = "EEE 402 (2026) – Final Project | Group 13 (G2)"
COVER_BAND = ["EEE 402 – Artificial Intelligence and Machine Learning Laboratory", "2026  ·  Lab Group G2  ·  Group 13", "Final Project Demonstration"]
FOOT_MID = "CT-based ADC vs SCC Classification   |   Presented by: 2106119"


def rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


TEAL, BLUE, GREEN, RED, ORANGE = "1485A4", "2683C6", "42BA97", "C00000", "F49100"
DARK, GREY, LIGHT, PALE = "262626", "595959", "E3F1F5", "F2F2F2"
FONT = "Arial"

prs = Presentation(TEMPLATE)
# drop every template slide, keep master + layouts
lst = prs.slides._sldIdLst
for sid in list(lst):
    prs.part.drop_rel(sid.rId)
    lst.remove(sid)

L_TITLE, L_ONLY = prs.slide_layouts[0], prs.slide_layouts[4]


# ---------------------------------------------------------------- helpers
def _runs(p, text, size, color, bold=False, italic=False):
    """'**x**' toggles bold."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        r = p.add_run()
        b = part.startswith("**")
        r.text = part[2:-2] if b else part
        f = r.font
        f.size, f.name, f.bold, f.italic = Pt(size), FONT, bold or b, italic
        f.color.rgb = rgb(color)


def _bullet(p, char="•", color=TEAL, level=0, size=16):
    pPr = p._p.get_or_add_pPr()
    indent = int(Pt(size) * 1.0)
    pPr.set("marL", str(indent * (level + 1) + (Inches(0.25) if level else 0)))
    pPr.set("indent", str(-indent))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(buClr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def text(slide, x, y, w, h, lines, size=16, color=DARK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False, gap=6, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        lvl = 0
        if isinstance(ln, tuple):
            ln, lvl = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        _runs(p, ln, size - 2 * lvl, color, bold, italic)
        if bullets:
            _bullet(p, "•" if lvl == 0 else "–", TEAL if lvl == 0 else GREY, lvl, size)
    return tb


def box(slide, x, y, w, h, fill, label="", size=14, color="FFFFFF", bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        line=None, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.5)
    else:
        s.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.06))
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = label if isinstance(label, list) else [label]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        sz, b = size, bold
        if isinstance(ln, tuple):
            ln, sz, b = ln
        _runs(p, ln, sz, color, b)
    return s


def arrow(slide, x1, y1, x2, y2, color=GREY, width=2.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def pic(slide, name, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(os.path.join(FIG, name), Inches(x), Inches(y), **kw)


def takeaway(slide, msg, y=6.2, h=0.6, x=0.6, w=12.1):
    b = box(slide, x, y, w, h, LIGHT, msg, size=15, color="0E5F75", bold=False, align=PP_ALIGN.LEFT)
    b.line.color.rgb = rgb(TEAL)
    b.line.width = Pt(1.25)
    b.text_frame.margin_left = Inches(0.2)
    return b


def table(slide, x, y, w, h, rows, widths, size=12, header=TEAL, zebra=True, highlight=None, bold_first=False):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for j, cw in enumerate(widths):
        t.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 or len(str(v)) > 14 else PP_ALIGN.CENTER
            if i == 0:
                c.fill.solid()
                c.fill.fore_color.rgb = rgb(header)
                _runs(p, str(v), size, "FFFFFF", True)
            else:
                c.fill.solid()
                hl = highlight is not None and i in (highlight if isinstance(highlight, (list, tuple)) else [highlight])
                c.fill.fore_color.rgb = rgb("D5ECF2" if hl else (PALE if zebra and i % 2 == 0 else "FFFFFF"))
                _runs(p, str(v), size, DARK, hl or (bold_first and j == 0))
    return t


_slide_no = [0]


def footer(slide):
    _slide_no[0] += 1
    lay = slide.slide_layout
    for ph in lay.placeholders:
        if ph.placeholder_format.idx in (10, 11, 12):
            slide.shapes.clone_placeholder(ph)
    for ph in slide.placeholders:
        idx = ph.placeholder_format.idx
        if idx == 10:
            ph.text_frame.text = FOOT_LEFT
        elif idx == 11:
            ph.text_frame.text = FOOT_MID
        elif idx == 12:
            p = ph.text_frame.paragraphs[0]
            fld = etree.SubElement(p._p, qn("a:fld"))
            fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
            fld.set("type", "slidenum")
            etree.SubElement(fld, qn("a:rPr")).set("lang", "en-US")
            etree.SubElement(fld, qn("a:t")).text = str(_slide_no[0])
        if idx in (10, 11):
            for r in ph.text_frame.paragraphs[0].runs:
                r.font.size = Pt(11)
                r.font.color.rgb = rgb("FFFFFF")


def new(title, sub=None):
    s = prs.slides.add_slide(L_ONLY)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.8)
    t.text_frame.text = title
    for r in t.text_frame.paragraphs[0].runs:
        r.font.size, r.font.name = Pt(32), FONT
    if sub:
        text(s, 0.6, 1.02, 12.1, 0.4, sub, size=15, color=GREY, italic=True)
    footer(s)
    return s


def chart_fmt(ch, size=12, legend=True):
    ch.font.size, ch.font.name = Pt(size), FONT
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)


def color_series(plot, cols):
    for s, c in zip(plot.series, cols):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(c)


def labels(plot, fmt="0.00", size=12, pos=XL_LABEL_POSITION.OUTSIDE_END):
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.font.size, dl.font.bold = Pt(size), True
    dl.position = pos


# ================================================================ 1 cover
s = prs.slides.add_slide(L_TITLE)
_slide_no[0] += 1
bg = s.shapes.add_picture(os.path.join(FIG, "cover_bg.png"), 0, Emu(1562266), prs.slide_width,
                          Emu(5968657 - 1562266))          # the white band between the two red bars
s.shapes._spTree.remove(bg._element)
s.shapes._spTree.insert(2, bg._element)                    # behind the title text
for sh in L_TITLE.shapes:
    if sh.name == "Rectangle 22":
        for para, new_txt in zip(sh.text_frame.paragraphs, COVER_BAND):
            para.runs[0].text = new_txt
            for r in para.runs[1:]:
                r.text = ""
t = s.shapes.title
t.left, t.top, t.width, t.height = Inches(2.8), Inches(2.05), Inches(7.73), Inches(1.7)
t.text_frame.text = "Telling Lung Adenocarcinoma from Squamous Cell Carcinoma on CT"
for r in t.text_frame.paragraphs[0].runs:
    r.font.size, r.font.name, r.font.bold = Pt(33), "Century Gothic", True
    r.font.color.rgb = rgb(TEAL)
sub = s.placeholders[1]
sub.left, sub.top, sub.width, sub.height = Inches(2.8), Inches(3.8), Inches(7.73), Inches(0.5)
sub.text_frame.text = "A radiomics model that still works on hospitals it has never seen"
for r in sub.text_frame.paragraphs[0].runs:
    r.font.size, r.font.name, r.font.italic = Pt(15), "Century Gothic", True
    r.font.color.rgb = rgb(GREY)
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.7), Inches(4.45), Inches(8.63), Inches(4.45))
ln.line.color.rgb, ln.line.width = rgb(RED), Pt(2)
text(s, 2.8, 4.6, 7.73, 0.45, "Rahib Mahasin  –  2106119", size=20, bold=True, align=PP_ALIGN.CENTER)
text(s, 2.8, 5.1, 7.73, 0.45, "Submitted by – Group 13, Lab Group G2", size=16, color=GREY, align=PP_ALIGN.CENTER)
for ph in list(s.placeholders):
    if ph.placeholder_format.idx in (11, 12):
        ph._element.getparent().remove(ph._element)

# ================================================================ 2 outline
s = new("Outline")
steps = [("1", "Problem", "What we want to solve and the hidden trap"),
         ("2", "Dataset", "4 hospitals + 2 locked test sets"),
         ("3", "Preprocessing", "Resampling, tumour outlines, features, ComBat"),
         ("4", "Model", "Feature blocks and classifier ensemble"),
         ("5", "Results", "Metrics, ablation, tuning, external test, comparison"),
         ("6", "Novelty", "What we did differently and why"),
         ("7", "Conclusion", "Limitations and future work")]
cols = [TEAL, BLUE, TEAL, BLUE, TEAL, BLUE, TEAL]
for i, (n, nm, d) in enumerate(steps):
    x = 0.55 + i * 1.75
    ch = box(s, x, 1.55, 1.85, 1.35, cols[i], [(n, 34, True)], shape=MSO_SHAPE.CHEVRON if i else
             MSO_SHAPE.PENTAGON)
    text(s, x, 3.1, 1.75, 0.5, nm, size=17, color=cols[i], bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.02, 3.62, 1.71, 1.3, d, size=14, color=GREY, align=PP_ALIGN.CENTER)
box(s, 1.2, 5.1, 10.9, 1.3, LIGHT,
    ["**In one sentence:** we built a CT-only model that tells ADC from SCC, and we tested it honestly – on "
     "hospitals it had never seen, with test labels locked away until the very end."],
    size=19, color="0E5F75", bold=False).line.color.rgb = rgb(TEAL)

# ================================================================ 3 problem
s = new("1. Problem Statement")
text(s, 0.6, 1.3, 7.2, 4.8, [
    "Non-small cell lung cancer has two main types: **adenocarcinoma (ADC)** and **squamous cell carcinoma "
    "(SCC)**. Together they make up most lung cancer cases.",
    "They are treated differently – some chemotherapy drugs and targeted therapies are suitable for only one of "
    "the two types, so doctors must know the type before treatment starts.",
    "Today the type is confirmed by a **biopsy** (needle or surgery). It is invasive, has risks, takes time, "
    "and sometimes the sample is too small or cannot be taken at all.",
    "But almost every lung cancer patient already has a **CT scan**. So our question is: can a computer tell "
    "ADC from SCC just by looking at the CT?"], size=17, bullets=True, gap=10)
box(s, 8.3, 1.45, 2.0, 0.9, BLUE, ["Chest CT", ("(3-D scan)", 12, False)])
arrow(s, 9.3, 2.35, 9.3, 2.75)
box(s, 8.3, 2.75, 2.0, 0.9, TEAL, ["Our model", ("(no biopsy)", 12, False)])
arrow(s, 9.3, 3.65, 9.3, 4.05)
box(s, 8.0, 4.05, 1.25, 0.8, GREEN, "ADC", size=16)
box(s, 9.35, 4.05, 1.25, 0.8, ORANGE, "SCC", size=16)
box(s, 10.85, 2.75, 1.85, 0.9, PALE, [("Needle biopsy", 13, True), ("invasive, slow", 11, False)], color=GREY,
    line="BFBFBF")
text(s, 10.85, 2.3, 1.85, 0.5, "✗", size=30, color=RED, bold=True, align=PP_ALIGN.CENTER)
takeaway(s, "**Goal:** input = chest CT + where the tumour is  →  output = probability of ADC vs SCC. "
            "And it must work on CT from hospitals the model has never seen before.", y=5.95, h=0.8)

# ================================================================ 4 shortcut
s = new("1.1 The Hidden Trap: the “Hospital Shortcut”")
text(s, 0.6, 1.3, 6.3, 4.6, [
    "Every hospital uses different scanners and settings, so its CT images carry a hidden “fingerprint” that "
    "a model can recognise very easily.",
    "Every hospital also has a different patient mix: **LUNG1 is 75% SCC**, while the other three hospitals "
    "are about **75% ADC**.",
    "So a model can learn “which hospital is this?” instead of “which cancer is this?” – and still get a "
    "high score when all hospitals are mixed together.",
    "We tested it: a model that only knows the **hospital name** and never looks at the image gets "
    "**AUC 0.71** on the normal metric!"], size=16, bullets=True, gap=9)
cd = CategoryChartData()
cd.categories = ["Hospital name only\n(no image!)", "Naive radiomics\nmodel"]
cd.add_series("Normal AUC (hospitals mixed)", (0.710, 0.727))
cd.add_series("Fair AUC (our metric)", (0.497, 0.642))
g = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.1), Inches(1.3), Inches(5.6), Inches(4.55), cd)
ch = g.chart
chart_fmt(ch, 13)
color_series(ch.plots[0], [RED, TEAL])
labels(ch.plots[0], size=13)
ch.plots[0].gap_width = 60
va = ch.value_axis
va.minimum_scale, va.maximum_scale, va.major_unit = 0.3, 0.8, 0.1
va.has_major_gridlines = False
va.tick_labels.font.size = Pt(12)
takeaway(s, "A high AUC on mixed hospitals can be fake. Our progress-presentation result (0.675) was partly "
            "affected by this – so for the final work we changed how we measure (Section 5).", y=6.05, h=0.75)

# ================================================================ 5 dataset
s = new("2. Dataset Description")
table(s, 0.6, 1.3, 7.3, 3.0, [
    ["Dataset (hospital)", "Country", "Patients", "ADC", "SCC"],
    ["LUNG1 (Maastro)", "Netherlands", "202", "51", "151"],
    ["Radiogenomics (Stanford)", "USA", "141", "112", "29"],
    ["Lung-PET-CT-Dx", "China", "301", "244", "57"],
    ["NLST (screening trial)", "USA", "397", "290", "107"],
    ["Total", "3 countries", "1041", "697", "344"]], [2.6, 1.4, 1.1, 1.1, 1.1], size=14, highlight=5)
text(s, 0.6, 4.45, 7.3, 1.7, [
    "**Data type:** 3-D chest CT scans (DICOM, about 100–600 slices per scan), all public (TCIA and NCI IDC).",
    "**Classes:** 2 – ADC = 1 (positive class), SCC = 0. The label comes from the pathology report, one label per "
    "patient.",
    "**Tumour location:** expert outlines (LUNG1, Radiogenomics) or expert boxes (Lung-PET-CT-Dx, NLST)."],
     size=14, bullets=True, gap=6)
cd = CategoryChartData()
cd.categories = ["LUNG1", "Radiogenomics", "Lung-PET-CT-Dx", "NLST"]
cd.add_series("ADC", (51, 112, 244, 290))
cd.add_series("SCC", (151, 29, 57, 107))
g = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED_100, Inches(8.2), Inches(1.25), Inches(4.6), Inches(4.3), cd)
ch = g.chart
chart_fmt(ch, 12)
ch.has_title = True
ch.chart_title.text_frame.text = "ADC / SCC share per hospital"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
color_series(ch.plots[0], [GREEN, ORANGE])
labels(ch.plots[0], fmt="0", size=12, pos=XL_LABEL_POSITION.CENTER)
ch.plots[0].gap_width = 50
ch.value_axis.has_major_gridlines = False
ch.value_axis.tick_labels.font.size = Pt(11)
ch.category_axis.tick_labels.font.size = Pt(11)
takeaway(s, "LUNG1 is “reversed” (mostly SCC) while the others are mostly ADC – exactly the situation where the "
            "hospital shortcut appears.", y=6.1, h=0.65)

# ================================================================ 6 locked tests
s = new("2.1 External Test Sets (Kept Locked)")
box(s, 0.6, 1.4, 5.1, 0.55, TEAL, "Development data – training & tuning", size=15)
for i, (nm, n) in enumerate([("LUNG1", 202), ("Radiogenomics", 141), ("Lung-PET-CT-Dx", 301), ("NLST", 397)]):
    box(s, 0.6 + (i % 2) * 2.6, 2.1 + (i // 2) * 1.0, 2.5, 0.85, "D5ECF2", [(nm, 14, True), (f"{n} patients", 12,
                                                                                          False)], color=DARK)
box(s, 6.05, 1.4, 0.22, 2.55, RED, "", shape=MSO_SHAPE.RECTANGLE)
text(s, 5.65, 4.0, 1.0, 0.5, "wall", size=12, color=RED, align=PP_ALIGN.CENTER, italic=True)
box(s, 6.6, 1.4, 6.1, 0.55, RED, "🔒  Locked test data – used only once, at the end", size=15)
box(s, 6.6, 2.1, 2.95, 1.85, "FBE5E5", [("Lung3", 16, True), ("Maastricht, surgery patients", 12, False),
                                        ("79 patients (44 ADC / 35 SCC)", 12, False),
                                        ("mostly 4–5 mm thick slices", 12, False)], color=DARK)
box(s, 9.75, 2.1, 2.95, 1.85, "FBE5E5", [("TCGA", 16, True), ("9 different US hospitals", 12, False),
                                         ("84 patients (48 ADC / 36 SCC)", 12, False),
                                         ("2/3 of scans are 5 mm thick", 12, False)], color=DARK)
text(s, 0.6, 4.45, 12.1, 1.6, [
    "These two sets were kept completely separate: we never trained, tuned or even looked at their labels until "
    "the very end.",
    "The tumour was found **automatically** (TotalSegmentator + MedSAM2) – no human drew anything, exactly like a "
    "new hospital would use the model.",
    "Before opening the answers we wrote the test plan down and saved a fingerprint (**SHA-256 hash**) of all "
    "predictions, so nothing could be changed afterwards. Each test was scored **only once**."],
     size=15, bullets=True, gap=6)
takeaway(s, "This is how a real new hospital would meet our model: no training, no tuning, no human help.",
         y=6.2, h=0.6)

# ================================================================ 7 pipeline
s = new("3. Full Pipeline at a Glance")
row1 = [("Chest CT (DICOM)", "hundreds of slices", BLUE), ("Convert & resample", "2 mm cubes, HU window", BLUE),
        ("Find tumour", "expert box or TotalSegmentator", TEAL), ("Outline tumour", "MedSAM2 (box prompt)", TEAL)]
row2 = [("Build rings", "0–4, 4–8, 8–12 mm", TEAL), ("Extract features", "texture + position + age/sex", GREEN),
        ("Remove hospital fingerprint", "ComBat", GREEN), ("Classifier ensemble", "4 models × 2 blocks", RED)]
BW, BH, GAPX = 2.6, 1.25, 0.55
Y1, Y2 = 1.45, 3.35
xs = [0.6 + i * (BW + GAPX) for i in range(4)]
for i, (a, b, c) in enumerate(row1):
    box(s, xs[i], Y1, BW, BH, c, [(a, 16, True), (b, 12, False)])
    if i < 3:
        arrow(s, xs[i] + BW, Y1 + BH / 2, xs[i + 1], Y1 + BH / 2)
arrow(s, xs[3] + BW / 2, Y1 + BH, xs[3] + BW / 2, Y2)          # down: outline -> rings
for i, (a, b, c) in enumerate(row2):                              # right-to-left
    j = 3 - i
    box(s, xs[j], Y2, BW, BH, c, [(a, 16, True), (b, 12, False)])
    if i < 3:
        arrow(s, xs[j], Y2 + BH / 2, xs[j - 1] + BW, Y2 + BH / 2)
Y3 = 5.2
box(s, xs[0], Y3, BW, 0.65, "FFFFFF", "ADC / SCC score", size=16, color=RED, line=RED)
arrow(s, xs[0] + BW / 2, Y2 + BH, xs[0] + BW / 2, Y3)             # down: classifier -> score
for i, (c, nm) in enumerate([(BLUE, "Image preparation"), (TEAL, "Tumour & regions"), (GREEN, "Features"),
                             (RED, "Model")]):
    box(s, 4.1 + i * 2.2, 5.37, 0.3, 0.3, c, "", shape=MSO_SHAPE.RECTANGLE)
    text(s, 4.45 + i * 2.2, 5.3, 1.8, 0.45, nm, size=13, color=GREY)
takeaway(s, "Every step runs automatically. The exact same frozen pipeline was later used on both locked external "
            "test sets.", y=6.15, h=0.65)

# ================================================================ 8 preprocessing images
s = new("3.1 Preprocessing: Making Every Scan Look the Same")
text(s, 0.6, 1.3, 6.2, 5.5, [
    "Raw CT comes as hundreds of separate DICOM slice files. We convert each scan into **one 3-D volume** "
    "(NIfTI format).",
    "Hospitals use different pixel sizes (0.6–1 mm) and slice gaps (1–5 mm). We **resample everything to "
    "2 × 2 × 2 mm cubes** (linear interpolation for CT, nearest-neighbour for masks), so a 1 cm tumour means the "
    "same thing everywhere.",
    "CT values are in Hounsfield Units (HU). We keep **−1024 to 200 HU** (air to soft tissue), which removes bone "
    "and metal that could confuse the texture features.",
    "Grey levels are grouped into **20 HU bins** before texture is calculated, so small noise has less effect."],
     size=17, bullets=True, gap=12)
pic(s, "ct_before_after.png", 7.0, 1.25, w=5.9)

# ================================================================ 9 masks
s = new("3.2 Preprocessing: Finding and Outlining the Tumour")
pic(s, "seg_gallery.png", 0.75, 1.2, w=11.8)
text(s, 0.6, 4.05, 7.4, 2.9, [
    "Texture features need an exact tumour outline, but drawing outlines by hand for 1000+ patients is not "
    "possible for us.",
    "We give **MedSAM2** (a medical version of Meta's “Segment Anything” model) one box around the tumour. It "
    "outlines the tumour and follows it slice by slice through the 3-D scan.",
    "For new hospitals where no box exists, **TotalSegmentator** finds the lung nodule first – so the whole "
    "process is automatic."], size=15, bullets=True, gap=7)
cd = CategoryChartData()
cd.categories = ["Tight box", "Loose box (+8 mm)"]
cd.add_series("Old threshold method", (0.82, 0.63))
cd.add_series("MedSAM2 (ours)", (0.85, 0.74))
g = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(8.2), Inches(3.95), Inches(4.5), Inches(2.95), cd)
ch = g.chart
chart_fmt(ch, 11)
ch.has_title = True
ch.chart_title.text_frame.text = "Overlap with expert outline (Dice, 1 = perfect)"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
color_series(ch.plots[0], [BLUE, GREEN])
labels(ch.plots[0], size=11)
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 0.4, 1.0
ch.value_axis.has_major_gridlines = False
ch.value_axis.visible = False

# ================================================================ 10 features
s = new("3.3 Features: Looking Inside and Around the Tumour")
pic(s, "shells_on_ct.png", 0.6, 1.3, w=5.3)
for i, (c, nm) in enumerate([(RED, "Tumour (GTV)"), (ORANGE, "Ring 0–4 mm"), (GREEN, "Ring 4–8 mm"),
                             (BLUE, "Ring 8–12 mm")]):
    box(s, 0.7 + (i % 2) * 2.2, 5.75 + (i // 2) * 0.45, 0.3, 0.3, c, "", shape=MSO_SHAPE.RECTANGLE)
    text(s, 1.07 + (i % 2) * 2.2, 5.68 + (i // 2) * 0.45, 1.9, 0.4, nm, size=14, color=GREY)
text(s, 6.15, 1.3, 6.6, 3.6, [
    "The tumour border carries information: ADC often grows along the air spaces with a fuzzy edge, while SCC is "
    "often more solid and central. So we look **inside and around** the tumour.",
    "We build **3 rings (shells)** outside the tumour – 0–4, 4–8 and 8–12 mm – and compute texture features in "
    "each ring with PyRadiomics.",
    "Texture families: first-order (brightness histogram), GLCM, GLRLM, GLSZM and NGTDM – they describe how grey "
    "levels repeat, form runs and form zones.",
    "**Position + clinical:** left/right lung, how central the tumour is, tumour volume, age and sex."],
     size=16, bullets=True, gap=8)
table(s, 6.15, 5.0, 6.6, 1.75, [["Feature block", "What it describes", "# features"],
                               ["Tumour texture (GTV)", "inside the tumour", "107"],
                               ["Rings (3 shells)", "tumour border & surroundings", "279"],
                               ["Position + clinical", "where it is, size, age, sex", "13"]],
      [2.3, 3.1, 1.2], size=14)

# ================================================================ 11 combat
s = new("3.4 Removing the Hospital Fingerprint (ComBat)")
pic(s, "pca_combat.png", 0.5, 1.25, w=7.6)
text(s, 8.3, 1.3, 4.45, 4.7, [
    "Even after resampling, every hospital's features are shifted (different scanners and reconstruction). In the "
    "left plot, LUNG1 (red) forms its own cloud.",
    "**ComBat** (first made for gene-data batch effects) removes each hospital's own average shift and spread, "
    "feature by feature.",
    "It is fitted **only on training patients** inside each fold. A new hospital is corrected using its images "
    "only – no labels needed.",
    "Check: a “hospital detector” found the hospital with AUC ≈ **0.9 before** ComBat and only ≈ **0.55–0.6 "
    "after**."], size=14, bullets=True, gap=7)
text(s, 0.6, 4.75, 7.4, 0.5, "Each dot = one patient (2-D PCA of the ring features), coloured by hospital.",
     size=12, color=GREY, italic=True, align=PP_ALIGN.CENTER)
takeaway(s, "After ComBat the hospitals overlap – the features now describe the tumour, not the scanner.",
         y=6.1, h=0.6)

# ================================================================ 12 architecture
s = new("4. Model Architecture")
for r, (nm, shape_txt, col, yy) in enumerate([("Ring features", "1041 × 279", TEAL, 1.45),
                                              ("Position + clinical", "1041 × 13", BLUE, 3.35)]):
    box(s, 0.6, yy, 1.9, 1.3, col, [(nm, 14, True), (shape_txt, 13, False)])
    arrow(s, 2.52, yy + 0.65, 2.88, yy + 0.65)
    box(s, 2.9, yy, 3.0, 1.3, "FFFFFF", [("Prep (training data only)", 13, True),
                                         ("impute → ComBat → z-score →", 11, False),
                                         ("rank → drop |r| > 0.9 → top-k", 11, False)], color=DARK, line=col)
    arrow(s, 5.92, yy + 0.65, 6.28, yy + 0.65)
    for j, m in enumerate(["Logistic reg. (L2)", "SVM (RBF)", "LightGBM", "Ridge (all)"]):
        box(s, 6.3, yy - 0.05 + j * 0.36, 2.0, 0.32, "D5ECF2", m, size=11, color=DARK, bold=False)
    arrow(s, 8.32, yy + 0.65, 8.68, yy + 0.65)
    box(s, 8.7, yy + 0.2, 1.3, 0.9, col, [("average", 13, True), ("block score", 11, False)],
        shape=MSO_SHAPE.OVAL)
    arrow(s, 10.02, yy + 0.65, 10.55, 2.95 if r == 0 else 3.45)
box(s, 10.55, 2.6, 2.15, 1.2, RED, [("Final score", 15, True), ("mean of 2 block scores", 11, False),
                                    ("> 0 → ADC,  < 0 → SCC", 11, False)])
table(s, 0.6, 5.0, 7.6, 1.25, [["Item", "Size / value"],
                               ["Input per patient", "1 vector: 279 ring + 13 position numbers"],
                               ["Features kept per model (top-k)", "10, 25 or 50 (tuned)"],
                               ["Output", "1 number per patient (ADC score)"]], [3.3, 4.3], size=12)
box(s, 8.5, 4.95, 4.2, 1.35, LIGHT, ["**Why not a big CNN?** With ~1000 patients, our deep models (CNN, foundation "
                                     "models) scored lower and learned the hospital (see 5.4)."],
    size=13, color="0E5F75", bold=False, align=PP_ALIGN.LEFT).line.color.rgb = rgb(TEAL)

# ================================================================ 13 training / tuning
s = new("4.1 Training Setup and Hyperparameter Tuning")
text(s, 0.6, 1.3, 5.8, 0.4, "Nested cross-validation (patient level)", size=15, bold=True, color=TEAL)
for f in range(5):
    box(s, 0.6 + f * 1.14, 1.8, 1.08, 0.55, ORANGE if f == 2 else BLUE, "Test" if f == 2 else f"Train",
        size=12, shape=MSO_SHAPE.RECTANGLE)
text(s, 0.6, 2.38, 5.7, 0.4, "Outer loop: 5 folds × 5 random seeds = 25 test runs", size=12, color=GREY,
     align=PP_ALIGN.CENTER)
for f in range(3):
    box(s, 1.2 + f * 1.55, 3.05, 1.45, 0.5, GREEN if f == 1 else "8EC3E6", "Inner val" if f == 1 else "Inner train",
        size=11, shape=MSO_SHAPE.RECTANGLE, color=DARK if f != 1 else "FFFFFF")
arrow(s, 2.3, 2.75, 2.3, 3.03)
text(s, 0.6, 3.6, 5.7, 0.4, "Inner loop (3 folds): picks hyperparameters", size=12, color=GREY,
     align=PP_ALIGN.CENTER)
text(s, 0.6, 4.15, 5.8, 2.2, [
    "Splits keep every hospital × class group in every fold.",
    "The test fold is **never** used for tuning or for ComBat.",
    "Sample weights make every hospital count equally and each 50/50 ADC/SCC (fixes class imbalance too)."],
     size=14, bullets=True, gap=5)
table(s, 6.7, 1.35, 6.0, 3.0, [["Model", "Hyperparameters tried", "Most often chosen"],
                               ["Logistic reg. (L2)", "C = 0.01 / 0.1 / 1; k = 10 / 25 / 50", "C = 0.01"],
                               ["SVM (RBF)", "C = 0.1 / 1 / 10; k", "C = 0.1"],
                               ["LightGBM", "leaves = 4 / 8; k (200 trees, lr 0.03)", "4 leaves"],
                               ["Ridge (all features)", "C = 0.0001 … 0.01", "C = 0.003–0.01"],
                               ["Selection metric", "fair (site-balanced) AUC", "–"]], [1.8, 2.6, 1.6], size=12)
box(s, 6.7, 4.55, 6.0, 1.6, LIGHT, ["**What tuning told us:** the strongest regularisation (smallest C, fewest leaves) "
                                    "won almost every time. With noisy features from 4 hospitals, simple models "
                                    "generalise better than flexible ones."], size=14, color="0E5F75", bold=False,
    align=PP_ALIGN.LEFT).line.color.rgb = rgb(TEAL)
takeaway(s, "Tuning never touches the test patients – so the numbers on the next slides are not optimistic.",
         y=6.2, h=0.6)

# ================================================================ 14 fair measure
s = new("5. Results: How We Measure Fairly")
text(s, 0.6, 1.3, 5.9, 0.45, "① Site-balanced (fair) AUC", size=18, bold=True, color=TEAL)
hosp = [("LUNG1", RED), ("Radiogen.", BLUE), ("LPCD", GREEN), ("NLST", ORANGE)]
for i, (nm, c) in enumerate(hosp):
    x = 0.8 + i * 1.4
    box(s, x, 1.9, 1.2, 0.5, c, nm, size=12, shape=MSO_SHAPE.RECTANGLE)
    box(s, x, 2.45, 0.58, 0.5, GREEN, "ADC", size=10, shape=MSO_SHAPE.RECTANGLE)
    box(s, x + 0.62, 2.45, 0.58, 0.5, ORANGE, "SCC", size=10, shape=MSO_SHAPE.RECTANGLE)
text(s, 0.6, 3.15, 5.9, 3.0, [
    "Normal AUC mixes all patients; when hospitals have different ADC/SCC mixes, just knowing the hospital "
    "already “helps”.",
    "We re-weight patients (weight = 1 / number in that hospital-and-class group), so **every hospital counts "
    "equally** and inside each hospital ADC and SCC count equally.",
    "A model that only guesses the hospital now scores **exactly 0.50**."], size=14, bullets=True, gap=6)
text(s, 6.9, 1.3, 5.9, 0.45, "② Leave-one-hospital-out", size=18, bold=True, color=TEAL)
for r in range(4):
    for c in range(4):
        test = r == c
        box(s, 7.1 + c * 1.05, 1.9 + r * 0.42, 0.95, 0.36, "BFBFBF" if test else hosp[c][1],
            "test" if test else "", size=10, shape=MSO_SHAPE.RECTANGLE, color=DARK)
    text(s, 11.35, 1.88 + r * 0.42, 1.4, 0.4, f"run {r + 1}", size=11, color=GREY)
text(s, 6.9, 3.65, 5.9, 2.5, [
    "Train on 3 hospitals, test on the 4th, which the model has **never seen**. Repeat 4 times.",
    "This is the closest thing to “install the model in a new hospital”.",
    "Coloured = training hospitals, grey = unseen test hospital."], size=14, bullets=True, gap=6)
takeaway(s, "Every result in this section uses these two fair tests – never the normal mixed-hospital AUC.",
         y=6.2, h=0.6)

# ================================================================ 15 main results
s = new("5.1 Main Results")
for i, (num, cap, c) in enumerate([("0.722", "fair AUC on 4 hospitals\n95% CI 0.69 – 0.76", TEAL),
                                   ("0.714", "average AUC on an\nunseen hospital", BLUE),
                                   ("1041", "patients from 4 hospitals\nin 3 countries", GREEN)]):
    box(s, 0.6 + i * 2.35, 1.35, 2.2, 1.9, c, [(num, 34, True), (cap, 11, False)])
text(s, 0.6, 3.5, 6.9, 2.6, [
    "All four hospitals are clearly above random (0.5). The best is **Lung-PET-CT-Dx (0.85)**, the hardest is "
    "**LUNG1 (0.63)**.",
    "LUNG1 is hard because 78% of its patients are late-stage (stage III), where ADC and SCC tumours look alike "
    "(similar size: 44 vs 48 ml).",
    "Every number is averaged over 5 seeds × 5 folds, so it is stable – not one lucky split."],
     size=15, bullets=True, gap=8)
pic(s, "roc_per_site.png", 7.7, 1.2, h=5.0)
takeaway(s, "The model works in every hospital, with a fair metric that the hospital shortcut cannot fool.",
         y=6.25, h=0.55)

# ================================================================ 16 LOSO
s = new("5.2 Result on Each Unseen Hospital")
cd = CategoryChartData()
cd.categories = ["LUNG1", "Radiogenomics", "Lung-PET-CT-Dx", "NLST", "Average"]
cd.add_series("AUC on unseen hospital", (0.625, 0.704, 0.835, 0.692, 0.714))
g = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.3), Inches(7.2), Inches(4.7), cd)
ch = g.chart
chart_fmt(ch, 13, legend=False)
ch.has_title = False
pl = ch.plots[0]
pl.gap_width = 55
labels(pl, size=14)
for j, c in enumerate([RED, BLUE, GREEN, ORANGE, TEAL]):
    pt = pl.series[0].points[j]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(c)
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale, ch.value_axis.major_unit = 0.4, 0.9, 0.1
ch.value_axis.has_major_gridlines = False
ch.value_axis.tick_labels.font.size = Pt(12)
text(s, 8.1, 1.3, 4.6, 3.5, [
    "Each bar: the model was trained on the **other 3 hospitals only**, then tested on this one.",
    "Unseen-hospital scores (0.63–0.84) are close to the pooled score (0.72), so the model is **not just "
    "memorising hospitals**.",
    "More hospitals helped: the unseen-hospital average went from **0.68 with 3 hospitals to 0.71 with 4**.",
    "Random guessing would be 0.50 for every bar."], size=16, bullets=True, gap=10)
box(s, 8.1, 4.85, 4.6, 1.15, "FFFFFF", ["Pooled fair AUC **0.722**  vs  unseen-hospital average **0.714**",
                                        ("→ only 0.008 is lost when the hospital is new", 13, False)],
    size=14, color=DARK, bold=False, line=TEAL)
takeaway(s, "This is the most realistic number in the project: what happens when the model meets a new hospital.",
         y=6.2, h=0.6)

# ================================================================ 17 ablation
s = new("5.3 Ablation Study: What Mattered?")
cd = CategoryChartData()
blocks = [("CT-FM (foundation model)", 0.662), ("FMCIB (deep CNN features)", 0.675), ("Inner tumour ring", 0.698),
          ("Rings (3 shells)", 0.700), ("Tumour texture (GTV)", 0.702), ("Position + clinical", 0.709),
          ("All 7 blocks together", 0.720), ("Rings + Position (FINAL)", 0.722)]
cd.categories = [b for b, _ in blocks]
cd.add_series("Fair AUC", [v for _, v in blocks])
g = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.25), Inches(6.4), Inches(4.8), cd)
ch = g.chart
chart_fmt(ch, 12, legend=False)
ch.has_title = False
pl = ch.plots[0]
pl.gap_width = 45
labels(pl, fmt="0.000", size=12)
for j in range(len(blocks)):
    pt = pl.series[0].points[j]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(RED if j < 2 else (TEAL if j == len(blocks) - 1 else "8EC3E6"))
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 0.6, 0.75
ch.value_axis.has_major_gridlines = False
ch.value_axis.visible = False
ch.category_axis.tick_labels.font.size = Pt(12)
table(s, 7.1, 1.3, 5.6, 3.6, [["Change we tested", "Effect"],
                              ["Threshold masks → MedSAM2 masks", "0.714 → 0.722"],
                              ["3 → 4 hospitals (unseen-hospital avg.)", "0.68 → 0.71"],
                              ["Best single block → 2 blocks fused", "0.709 → 0.722"],
                              ["2 blocks → all 7 blocks", "0.722 → 0.720 (no gain)"],
                              ["ComBat off → on", "AUC about the same, hospital detector 0.9 → 0.6"]],
      [3.2, 2.4], size=12)
takeaway(s, "Good tumour outlines + simple, well-chosen features mattered more than bigger deep models "
            "(red bars).", y=6.15, h=0.6)

# ================================================================ 18 what did not work
s = new("5.4 What Did NOT Work (and Why That Is Useful)")
table(s, 0.6, 1.35, 12.1, 3.9, [
    ["Idea we tried", "Result (fair AUC)", "What we learned"],
    ["Attention-MIL on small local patches (progress presentation)", "≈ 0.60 (LUNG1)",
     "Attention did not beat simple mean pooling; too little data"],
    ["Frozen deep CNN features (FMCIB)", "0.675", "Good, but weaker than handcrafted ring features"],
    ["CT foundation model (CT-FM)", "0.662", "Kept remembering the hospital, even after ComBat (detector 0.77)"],
    ["Fine-tuned CNN + “hospital-confusing” adversarial loss", "0.606",
     "Overfitted; learned the hospital even more (detector 0.82)"],
    ["Fusing all 7 feature blocks", "0.720", "No gain over the simple 2-block model"]],
      [4.6, 2.2, 5.3], size=15)
for i in range(5):
    text(s, 0.25, 2.02 + i * 0.66, 0.4, 0.5, "✗", size=20, color=RED, bold=True)
takeaway(s, "Bigger models learned the hospital more than the cancer. Reporting negative results honestly is part "
            "of the work – it tells the next person what not to try.", y=5.6, h=0.8)

# ================================================================ 19 journey
s = new("5.5 Project Journey: From Progress Presentation to Final")
cd = CategoryChartData()
cd.categories = ["Attention-MIL\n(LUNG1)", "CNN + attention\n(LUNG1)", "2 cohorts merged\n(normal AUC)",
                 "Fair metric +\nComBat (2 hosp.)", "3 hospitals", "4 hospitals", "+ MedSAM2\n(FINAL)",
                 "Lung3 locked\ntest"]
cd.add_series("AUC", (0.604, 0.615, 0.675, 0.681, 0.718, 0.714, 0.722, 0.686))
g = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.6), Inches(1.3), Inches(12.1), Inches(4.3), cd)
ch = g.chart
chart_fmt(ch, 11, legend=False)
ch.has_title = False
pl = ch.plots[0]
labels(pl, size=13, pos=XL_LABEL_POSITION.ABOVE)
ser = pl.series[0]
ser.format.line.color.rgb = rgb(TEAL)
ser.format.line.width = Pt(3)
ser.smooth = False
ser.marker.size = 10
ser.marker.format.fill.solid()
ser.marker.format.fill.fore_color.rgb = rgb(TEAL)
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 0.55, 0.76
ch.value_axis.has_major_gridlines = False
ch.value_axis.tick_labels.font.size = Pt(11)
ch.category_axis.tick_labels.font.size = Pt(11)
box(s, 4.35, 1.3, 0.05, 3.6, RED, "", shape=MSO_SHAPE.RECTANGLE)
text(s, 1.2, 1.25, 3.1, 0.4, "◀ progress presentation", size=12, color=RED, bold=True, align=PP_ALIGN.RIGHT)
text(s, 4.5, 1.25, 3.5, 0.4, "after progress presentation ▶", size=12, color=TEAL, bold=True)
takeaway(s, "The biggest jumps came from **more hospitals, a fair metric and better tumour outlines** – not from more "
            "complex networks. The 0.675 before the red line was partly inflated by the hospital shortcut.",
         y=5.85, h=0.85)

# ================================================================ 20 external
s = new("5.6 Locked External Test on New Hospitals")
pic(s, "external_ci.png", 0.45, 1.25, w=7.0)
text(s, 7.75, 1.25, 5.0, 4.6, [
    "**Lung3: AUC 0.686** (sensitivity 0.64, specificity 0.69) – very close to our unseen-hospital CV (0.71). "
    "The model transferred.",
    "**TCGA: AUC 0.585** – lower. In TCGA each of the 9 hospitals sent **only ADC or only SCC** patients, so the "
    "test itself is mixed up with hospital effects (the trap from slide 4).",
    "**Thin-slice CT (≤ 2.5 mm) worked best** in both sets (0.90 and 0.73). Thick 5 mm slices blur the texture.",
    "TCGA tumours were also smaller (16 vs 29 ml), so automatic detection was harder."],
     size=16, bullets=True, gap=9)
box(s, 0.6, 5.9, 12.1, 0.9, "FBE5E5",
    ["🔒 **Honesty check:** the test plan was written first, predictions were locked with a SHA-256 hash before the "
     "labels were opened, and each test was scored only once. Error bars = 95% bootstrap confidence interval."],
    size=15, color="7A0000", bold=False, align=PP_ALIGN.LEFT).line.color.rgb = rgb(RED)

# ================================================================ 21 comparison
s = new("5.7 Comparison with Existing Works")
table(s, 0.45, 1.2, 12.45, 5.0, [
    ["Study", "Data", "How it was tested", "AUC", "Remarks"],
    ["Zhu et al. 2018, Eur Radiol", "129 pts, 1 hospital", "Split inside the same hospital", "0.91",
     "Small, single hospital, no external test"],
    ["Pasini et al. 2023, Diagnostics", "466 pts incl. LUNG1 + Radiogenomics", "Random 80/20 split",
     "Acc 0.77 → 0.59", "Accuracy fell after ComBat – same hospital effect we found"],
    ["Yang et al. 2021, Front Oncol", "645 pts, 3 centres", "Train 1 centre → test others", "0.54–0.64",
     "0.78 only when centres were mixed randomly"],
    ["Chaunzwa et al. 2021, Sci Rep", "311 pts, 1 hospital (MGH)", "External: Lung3 (49 pts)", "0.71 → 0.60",
     "Deep CNN; same Lung3 test set as ours"],
    ["Chen et al. 2023, Radiol Med", "402 pts, TCIA", "Internal / external (78 pts)", "0.84 → 0.73",
     "Multi-task CNN; clear drop on external"],
    ["Song et al. 2023, Med Phys", "868 pts, 8 TCIA sets", "Internal + TCGA (97) + Lung3 (71)", "0.82 / 0.82 / 0.80",
     "Best of 130 models reported on these test sets; expert-drawn masks; hospitals mixed randomly"],
    ["Ours (this project)", "1041 pts, 4 hospitals", "Unseen hospital + 2 locked tests",
     "0.72 / 0.71 / Lung3 0.69 / TCGA 0.59", "Fair metric, fully automatic masks, pre-registered locked tests"]],
      [2.45, 2.35, 2.35, 1.75, 3.55], size=10.5, highlight=7, bold_first=True)
tk = takeaway(s, "Scores above 0.8 mostly come from one hospital, random splits or picking the best model on the "
                 "test set; on a new hospital most papers get 0.54–0.73. On the same Lung3 set we beat Chaunzwa "
                 "(0.69 vs 0.60); Song reports higher (0.80) but with expert masks and no locked test.",
              y=6.3, h=0.68)
for r in tk.text_frame.paragraphs[0].runs:
    r.font.size = Pt(13)

# ================================================================ 22 novelty
s = new("6. Novelty: What We Did Differently")
cards = [
    ("1", "Fair metric + fingerprint removal", TEAL,
     "**What:** site-balanced AUC and ComBat inside every fold.",
     "**Why:** mixed-hospital AUC can be fooled by the hospital shortcut (0.71 with no image!).",
     "**Result:** honest score rose from 0.64 (naive) to 0.72."),
    ("2", "Fully automatic tumour outlining", BLUE,
     "**What:** TotalSegmentator finds the tumour, MedSAM2 outlines it from a single box.",
     "**Why:** hand-drawing is impossible at a new hospital.",
     "**Result:** Dice 0.85 vs experts; masks helped AUC 0.714 → 0.722."),
    ("3", "4 hospitals, 1041 patients", GREEN,
     "**What:** the largest public mixed-hospital set we know of for this task (3 countries).",
     "**Why:** a model trained at one hospital does not travel.",
     "**Result:** 0.71 on hospitals it never saw (0.68 with 3)."),
    ("4", "Locked, pre-registered external tests", RED,
     "**What:** plan written first, predictions hashed, scored once.",
     "**Why:** no chance to “tune on the test set”.",
     "**Result:** honest external numbers: Lung3 0.69, TCGA 0.59.")]
for i, (n, ttl, c, a, b, r) in enumerate(cards):
    x = 0.6 + i * 3.07
    box(s, x, 1.3, 2.9, 0.95, c, [(f"{n}.  {ttl}", 15, True)])
    card = box(s, x, 2.3, 2.9, 3.6, "FFFFFF", "", line=c)
    text(s, x + 0.1, 2.42, 2.7, 3.45, [a, b, r], size=16, gap=14)
takeaway(s, "In short: most papers try to raise the AUC. We first made sure the AUC is honest, then raised it.",
         y=6.1, h=0.6)

# ================================================================ 23 conclusion
s = new("7. Conclusion, Limitations and Future Work")
box(s, 0.6, 1.3, 12.1, 1.0, LIGHT,
    ["**Conclusion:** a fully automatic CT model separates ADC from SCC with a fair AUC of 0.72 on 4 hospitals, "
     "0.71 on unseen hospitals and 0.69 on a locked external test (Lung3)."],
    size=17, color="0E5F75", bold=False, align=PP_ALIGN.LEFT).line.color.rgb = rgb(TEAL)
box(s, 0.6, 2.45, 5.9, 0.5, RED, "Limitations", size=17)
text(s, 0.6, 3.05, 5.9, 3.0, [
    "Thick 5 mm CT slices hurt a lot (TCGA thick slices: 0.51).",
    "LUNG1 (late-stage patients) is still hard at about 0.63.",
    "External test sets are small (about 80 patients each), so their confidence intervals are wide.",
    "About 1 in 5 automatic detections picks the wrong nodule."], size=17, bullets=True, gap=12)
box(s, 6.8, 2.45, 5.9, 0.5, GREEN, "Future work", size=17)
text(s, 6.8, 3.05, 5.9, 3.0, [
    "Train with simulated thick slices so the model does not depend on slice thickness.",
    "Add PET scans, which carry metabolic information that CT does not.",
    "Test on a bigger external set with both classes from every hospital.",
    "Let a radiologist quickly confirm the detected nodule before prediction."], size=17, bullets=True, gap=12)

# ================================================================ 24 references
s = new("8. References")
text(s, 0.6, 1.2, 12.1, 5.75, [
    "Zhu X. et al., “Radiomic signature as a diagnostic factor for histologic subtype classification of NSCLC,” "
    "Eur Radiol, 2018.",
    "Pasini G. et al., “Phenotyping the histopathological subtypes of NSCLC: how beneficial is radiomics?,” "
    "Diagnostics, 2023.",
    "Yang F. et al., “Machine learning for histologic subtype classification of NSCLC: a retrospective multicenter "
    "radiomics study,” Front Oncol, 2021.",
    "Chaunzwa T. et al., “Deep learning classification of lung cancer histology using CT images,” Sci Rep, 2021.",
    "Chen et al., multi-task CNN for NSCLC subtype and stage, Radiol Med, 2023.",
    "Song F. et al., “Radiomics feature analysis and model research for predicting histopathological subtypes of "
    "NSCLC on CT images: a multi-dataset study,” Med Phys, 2023.",
    "Ma J. et al., “MedSAM2: Segment anything in 3D medical images and videos,” 2025.",
    "Wasserthal J. et al., “TotalSegmentator: robust segmentation of 104 anatomical structures in CT,” "
    "Radiol AI, 2023.",
    "Johnson W. E. et al., “Adjusting batch effects in microarray data using empirical Bayes methods (ComBat),” "
    "Biostatistics, 2007.",
    "van Griethuysen J. et al., “Computational radiomics system to decode the radiographic phenotype "
    "(PyRadiomics),” Cancer Res, 2017.",
    "Datasets: TCIA NSCLC-Radiomics (LUNG1), NSCLC-Radiogenomics, Lung-PET-CT-Dx, NSCLC-Radiomics-Genomics "
    "(Lung3), TCGA-LUAD/LUSC; NCI IDC NLST."], size=15, bullets=True, gap=7)

# ================================================================ 25 thank you
s = new("")
s.shapes.title._element.getparent().remove(s.shapes.title._element)
text(s, 0.6, 2.2, 12.1, 1.3, "Thank You!", size=60, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
text(s, 0.6, 3.6, 12.1, 0.6, "Questions?", size=28, color=GREY, align=PP_ALIGN.CENTER)
ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.7), Inches(4.45), Inches(8.63), Inches(4.45))
ln.line.color.rgb, ln.line.width = rgb(RED), Pt(2)
text(s, 0.6, 4.65, 12.1, 0.5, "Rahib Mahasin  –  2106119  |  Group 13, Lab Group G2", size=16, color=GREY,
     align=PP_ALIGN.CENTER)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
