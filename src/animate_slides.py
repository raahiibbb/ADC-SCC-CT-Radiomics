"""Add simple, modern animations to the finished deck (in place).

- every slide: soft fade transition
- content (not titles/footers): fade + small upward float, played automatically
  one after another; bullet lists build paragraph by paragraph; runs of small
  shapes (flowchart boxes, arrows) cascade; charts wipe in (bars "grow").
Idempotent: existing transitions/animations are replaced.
Usage (.venv-phase10):  python src/animate_slides.py [pptx]
"""
from __future__ import annotations

import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT = os.path.join(os.path.dirname(HERE), "presentation", "ADC_SCC_Final_Presentation.pptx")

DUR, CHART_DUR, STAGGER, PARA_STAGGER, OVERLAP = 250, 450, 60, 120, 60
BUDGET = 900                      # ms: every slide finishes its entrance within this


def E(tag, parent=None, **attrs):
    el = etree.Element(qn(tag)) if parent is None else etree.SubElement(parent, qn(tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


class Ids:
    def __init__(self):
        self.n = 2

    def __call__(self):
        self.n += 1
        return str(self.n)


def _target(parent, spid, para=None):
    tgt = E("p:tgtEl", parent)
    sp = E("p:spTgt", tgt, spid=spid)
    if para is not None:
        tx = E("p:txEl", sp)
        E("p:pRg", tx, st=para, end=para)


def _visible(parent, ids, spid, para):
    s = E("p:set", parent)
    b = E("p:cBhvr", s)
    c = E("p:cTn", b, id=ids(), dur="1", fill="hold")
    E("p:cond", E("p:stCondLst", c), delay="0")
    _target(b, spid, para)
    E("p:attrName", E("p:attrNameLst", b)).text = "style.visibility"
    E("p:strVal", E("p:to", s), val="visible")


def _anim_y(parent, ids, spid, para, dur):
    for attr, start in (("ppt_x", "#ppt_x"), ("ppt_y", "#ppt_y+0.03")):
        a = E("p:anim", parent, calcmode="lin", valueType="num")
        b = E("p:cBhvr", a)
        E("p:cTn", b, id=ids(), dur=dur, fill="hold")
        _target(b, spid, para)
        E("p:attrName", E("p:attrNameLst", b)).text = attr
        tl = E("p:tavLst", a)
        for tm, val in (("0", start), ("100000", "#" + attr)):
            E("p:strVal", E("p:val", E("p:tav", tl, tm=tm)), val=val)


def effect(parent, ids, spid, kind, delay, first, para=None):
    """kind: float | fade | wipe-up | wipe-right."""
    preset = {"float": ("42", "0"), "fade": ("10", "0"), "wipe-up": ("22", "4"), "wipe-right": ("22", "8")}[kind]
    par = E("p:par", parent)
    c = E("p:cTn", par, id=ids(), presetID=preset[0], presetClass="entr", presetSubtype=preset[1], fill="hold",
          grpId="0", nodeType="afterEffect" if first else "withEffect")
    E("p:cond", E("p:stCondLst", c), delay=int(delay))
    kids = E("p:childTnLst", c)
    _visible(kids, ids, spid, para)
    dur = CHART_DUR if kind.startswith("wipe") else DUR
    filt = {"float": "fade", "fade": "fade", "wipe-up": "wipe(down)", "wipe-right": "wipe(left)"}[kind]
    ae = E("p:animEffect", kids, transition="in", filter=filt)
    b = E("p:cBhvr", ae)
    E("p:cTn", b, id=ids(), dur=dur)
    _target(b, spid, para)
    if kind == "float":
        _anim_y(kids, ids, spid, para, dur)
    return dur


def classify(sh):
    """-> ('bullets', n_paras) | ('chart', kind) | ('major', None) | ('minor', None)."""
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "major", None
    if sh.has_chart:
        t = sh.chart.chart_type
        horiz = t in (XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.BAR_STACKED) or "LINE" in str(t)
        return "chart", "wipe-right" if horiz else "wipe-up"
    if sh.has_table:
        return "major", None
    if sh.is_placeholder:
        return "major", None
    if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and sh.has_text_frame:
        paras = [p for p in sh.text_frame.paragraphs if p.text.strip()]
        if len(paras) > 6:                       # long lists (e.g. references): one quick fade
            return "major", None
        if len(paras) >= 2:
            return "bullets", len(sh.text_frame.paragraphs)
        if sh.width > Inches(6):
            return "major", None
    return "minor", None


def steps_for(slide, skip_ids, animate_title=False):
    shapes = [sh for sh in slide.shapes if sh.shape_id not in skip_ids and
              (not sh.is_placeholder or (animate_title and sh.placeholder_format.idx in (0, 1)))]
    steps, run = [], []
    for sh in shapes:
        cls, arg = classify(sh)
        if cls == "minor":
            run.append(sh)
            continue
        if run:
            steps.append(("run", run))
            run = []
        steps.append((cls, (sh, arg)))
    if run:
        steps.append(("run", run))
    return shapes, steps


def animate(slide, skip_ids=(), animate_title=False):
    sld = slide._element
    for tag in ("p:transition", "p:timing"):
        for el in sld.findall(qn(tag)):
            sld.remove(el)
    # soft fade transition (must sit after clrMapOvr)
    tr = E("p:transition", spd="fast")
    E("p:fade", tr)
    anchor = sld.find(qn("p:clrMapOvr"))
    (anchor.addnext if anchor is not None else sld.find(qn("p:cSld")).addnext)(tr)

    shapes, steps = steps_for(slide, set(skip_ids), animate_title)
    if not steps:
        return 0
    ids = Ids()
    timing = E("p:timing")
    root = E("p:cTn", E("p:par", E("p:tnLst", timing)), id="1", dur="indefinite", restart="never", nodeType="tmRoot")
    seq = E("p:seq", E("p:childTnLst", root), concurrent="1", nextAc="seek")
    main = E("p:cTn", seq, id="2", dur="indefinite", nodeType="mainSeq")
    click = E("p:par", E("p:childTnLst", main))
    cc = E("p:cTn", click, id=ids(), fill="hold")
    sc = E("p:stCondLst", cc)
    E("p:cond", sc, delay="indefinite")
    E("p:tn", E("p:cond", sc, evt="onBegin", delay="0"), val="2")      # start automatically with the slide
    steps_parent = E("p:childTnLst", cc)
    for tag, evt in (("p:prevCondLst", "onPrev"), ("p:nextCondLst", "onNext")):
        E("p:sldTgt", E("p:tgtEl", E("p:cond", E(tag, seq), evt=evt, delay="0")))

    # 1) plan every effect with its natural start time
    plan, bld_items, t = [], [], 0.0
    for kind, payload in steps:
        if kind == "run":
            stag = min(STAGGER, 500 / max(len(payload), 1))
            end = 0
            for i, sh in enumerate(payload):
                is_line = sh.element.tag.endswith("cxnSp")
                plan.append((str(sh.shape_id), "fade" if is_line else "float", None, t + i * stag))
                end = i * stag + DUR
                if sh.has_text_frame and sh.text_frame.text.strip():
                    bld_items.append(("p", str(sh.shape_id), "animBg"))
            t += end - OVERLAP
        elif kind == "bullets":
            sh, n = payload
            paras = [i for i, p in enumerate(sh.text_frame.paragraphs) if p.text.strip()]
            for j, pi in enumerate(paras):
                plan.append((str(sh.shape_id), "float", pi, t + j * PARA_STAGGER))
            bld_items.append(("p", str(sh.shape_id), "build"))
            t += (len(paras) - 1) * PARA_STAGGER + DUR - OVERLAP
        else:
            sh, arg = payload
            k = arg if kind == "chart" else "float"
            plan.append((str(sh.shape_id), k, None, t))
            if sh.has_chart or sh.has_table:
                bld_items.append(("g", str(sh.shape_id), None))
            elif sh.has_text_frame and sh.text_frame.text.strip():
                bld_items.append(("p", str(sh.shape_id), "animBg"))
            t += (CHART_DUR if kind == "chart" else DUR) - OVERLAP
    # 2) squeeze start times so every slide finishes within the same budget
    f = min([1.0] + [(BUDGET - (CHART_DUR if k.startswith("wipe") else DUR)) / s for _, k, _, s in plan if s > 0])
    # 3) one automatic group: first effect "after previous", the rest "with previous" + offset
    sp = E("p:par", steps_parent)
    spc = E("p:cTn", sp, id=ids(), fill="hold")
    E("p:cond", E("p:stCondLst", spc), delay="0")
    kids = E("p:childTnLst", spc)
    for i, (spid, k, para, start) in enumerate(plan):
        effect(kids, ids, spid, k, start * f, i == 0, para=para)
    bld = E("p:bldLst")
    for typ, spid, how in bld_items:
        if typ == "g":
            E("p:bldAsOne", E("p:bldGraphic", bld, spid=spid, grpId="0"))
        elif how == "build":
            E("p:bldP", bld, spid=spid, grpId="0", build="p")
        else:
            E("p:bldP", bld, spid=spid, grpId="0", animBg="1")
    if len(bld):
        timing.append(bld)
    tr.addnext(timing)
    return len(steps)


def main(path):
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides):
        skip = set()
        if i == 0:                       # cover: keep the background art static
            skip = {sh.shape_id for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE}
        n = animate(slide, skip, animate_title=(i == 0))
        print(f"slide {i + 1}: {n} animation steps")
    prs.save(path)
    print("saved", path)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else DEFAULT)
